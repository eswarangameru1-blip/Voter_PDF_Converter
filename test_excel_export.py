import os
import unittest
import openpyxl
from excel_export import (
    create_workbook,
    style_header,
    style_rows,
    auto_fit_columns,
    create_table,
    freeze_header,
    save_workbook,
    export_to_excel
)

class TestExcelExport(unittest.TestCase):
    def setUp(self):
        # Sample mock data
        self.mock_voters = [
            {"S.No": 1, "VoterID Number": "ABC1234567", "Elector's Name": "Voter One", "Guardian Name": "Guardian One", "House Number": "1-10", "Sex": "Male", "Age": 25},
            {"S.No": 2, "VoterID Number": "XYZ9876543", "Elector's Name": "Voter Two", "Guardian Name": "Guardian Two", "House Number": "2-20", "Sex": "Female", "Age": 30},
            {"S.No": 3, "VoterID Number": "ABC1234567", "Elector's Name": "Voter Three (Dup)", "Guardian Name": "Guardian Three", "House Number": "3-30", "Sex": "Male", "Age": 45}, # Duplicate EPIC
            {"S.No": "invalid", "VoterID Number": "MNO4561234", "Elector's Name": "Voter Four", "Guardian Name": "Guardian Four", "House Number": "4-40", "Sex": "Female", "Age": 22}, # Invalid S.No
            {"S.No": 5, "VoterID Number": "JKL7890123", "Elector's Name": "Voter Five", "Guardian Name": "Guardian Five", "House Number": "5-50", "Sex": "Male", "Age": "invalid"}, # Invalid Age
            {"S.No": 6, "VoterID Number": "", "Elector's Name": "Voter Six", "Guardian Name": "Guardian Six", "House Number": "6-60", "Sex": "Female", "Age": 33}, # Empty EPIC
        ]
        self.test_pdf = "test_export_voters.pdf"
        self.expected_file = "output/test_export_voters_voter_details.xlsx"

    def tearDown(self):
        # Clean up generated Excel files
        if os.path.exists(self.expected_file):
            os.remove(self.expected_file)
        
        # Clean up any incremented files too
        for i in range(1, 5):
            inc_file = f"output/test_export_voters_voter_details_{i}.xlsx"
            if os.path.exists(inc_file):
                os.remove(inc_file)

    def test_workbook_creation(self):
        """Verifies workbook creation helper sets properties properly."""
        wb = create_workbook()
        self.assertIsNotNone(wb)
        self.assertEqual(wb.properties.title, "Voter Details")
        self.assertEqual(wb.properties.creator, "Voter PDF Converter")
        self.assertEqual(wb.properties.subject, "Election Commission Voter Records")
        
        # Check Custom properties
        custom_props = {p.name: p.value for p in wb.custom_doc_props.props}
        self.assertIn("Company", custom_props)
        self.assertEqual(custom_props["Company"], "Voter PDF Converter")

    def test_header_formatting_and_column_order(self):
        """Verifies header formatting, columns appending, and column order."""
        wb = openpyxl.Workbook()
        ws = wb.active
        columns = ["S.No", "VoterID Number", "Elector's Name", "Guardian Name", "House Number", "Sex", "Age"]
        
        style_header(ws, columns)
        
        # Verify row count has header row
        self.assertEqual(ws.max_row, 1)
        
        # Verify column values and order
        row_values = [ws.cell(row=1, column=c).value for c in range(1, 8)]
        self.assertEqual(row_values, columns)
        
        # Verify header styling (font color = FFFFFF (white), fill color = 1F4E78)
        cell = ws.cell(row=1, column=1)
        self.assertIn(cell.font.color.value, ["FFFFFF", "00FFFFFF"])
        self.assertEqual(cell.font.bold, True)
        self.assertIn(cell.fill.start_color.value, ["1F4E78", "001F4E78"])
        self.assertEqual(cell.alignment.horizontal, "center")

    def test_row_count_and_duplicate_removal(self):
        """Verifies validation and duplicate EPIC filtering during export."""
        res = export_to_excel(self.mock_voters, self.test_pdf)
        
        self.assertEqual(res["status"], "success")
        self.assertEqual(res["records"], 2)  # Voter 1 and 2 are fully valid.
        self.assertEqual(res["duplicates_removed"], 1) # Voter 3 has duplicate EPIC "ABC1234567"
        self.assertEqual(res["invalid_records"], 3) # Voter 4 (s_no), 5 (age), 6 (empty EPIC)

        # Load generated file to verify contents
        wb = openpyxl.load_workbook(res["file"])
        ws = wb.active
        
        # Header + 2 data rows = 3 rows total
        self.assertEqual(ws.max_row, 3)
        
        # Check data values
        self.assertEqual(ws.cell(row=2, column=1).value, 1)
        self.assertEqual(ws.cell(row=2, column=2).value, "ABC1234567")
        self.assertEqual(ws.cell(row=3, column=1).value, 2)
        self.assertEqual(ws.cell(row=3, column=2).value, "XYZ9876543")

    def test_file_creation_and_increments(self):
        """Verifies workbook is successfully created on disk and increments filenames correctly."""
        res1 = export_to_excel(self.mock_voters, self.test_pdf)
        file1 = res1["file"]
        self.assertTrue(os.path.exists(file1))
        self.assertEqual(file1, "output/test_export_voters_voter_details.xlsx")
        
        # Export again with same name to test filename incrementing
        res2 = export_to_excel(self.mock_voters, self.test_pdf)
        file2 = res2["file"]
        self.assertTrue(os.path.exists(file2))
        self.assertEqual(file2, "output/test_export_voters_voter_details_1.xlsx")

if __name__ == '__main__':
    unittest.main()
