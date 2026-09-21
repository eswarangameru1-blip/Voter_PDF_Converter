import os
import time
import logging
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.packaging.custom import StringProperty
from typing import List, Dict, Any

logger = logging.getLogger(__name__)

def create_workbook() -> openpyxl.Workbook:
    """
    Helper: Creates an Excel workbook with correct metadata properties.
    """
    wb = openpyxl.Workbook()
    wb.properties.title = "Voter Details"
    wb.properties.creator = "Voter PDF Converter"
    wb.properties.subject = "Election Commission Voter Records"
    
    # Custom extended properties
    wb.custom_doc_props.append(StringProperty(name='Company', value='Voter PDF Converter'))
    return wb

def style_header(ws, columns: List[str]):
    """
    Helper: Appends columns and styles the header row.
    Header styling: Bold, White font, Dark Blue fill (#1F4E78), Center alignment, Thin borders.
    """
    ws.append(columns)
    
    font_header = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    fill_header = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
    align_center = Alignment(horizontal='center', vertical='center')
    thin_side = Side(style='thin', color='D9D9D9')
    border_all = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    
    for col_idx in range(1, len(columns) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = border_all

def style_rows(ws, start_row: int, voters: List[Dict[str, Any]]):
    """
    Helper: Writes voter data rows and applies styling.
    Row styling: Alternating row colors, thin borders, vertical center alignment.
    """
    font_body = Font(name='Calibri', size=11, bold=False)
    fill_zebra = PatternFill(start_color='F2F4F7', end_color='F2F4F7', fill_type='solid')
    fill_review = PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid')
    
    align_center = Alignment(horizontal='center', vertical='center')
    align_left = Alignment(horizontal='left', vertical='center')
    thin_side = Side(style='thin', color='D9D9D9')
    border_all = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    
    for offset, voter in enumerate(voters):
        row_idx = start_row + offset
        row_data = [
            voter.get("S.No"),
            voter.get("VoterID Number"),
            voter.get("Elector's Name"),
            voter.get("Guardian Name"),
            voter.get("House Number"),
            voter.get("Sex"),
            voter.get("Age")
        ]
        ws.append(row_data)
        
        is_review = voter.get("review_required", False)
        is_even = (row_idx % 2 == 0)
        row_fill = fill_review if is_review else (fill_zebra if is_even else None)
        
        for col_idx in range(1, 8):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = font_body
            cell.border = border_all
            
            if row_fill:
                cell.fill = row_fill
                
            if col_idx in [1, 2, 6, 7]:
                cell.alignment = align_center
            else:
                cell.alignment = align_left

def auto_fit_columns(ws, min_width: int = 15):
    """
    Helper: Adjusts columns width dynamically based on content with a minimum width.
    """
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 3, min_width)

def create_table(ws, max_row: int):
    """
    Helper: Converts a cell range into an Excel Table object with row striping enabled.
    """
    if max_row > 1:
        tab = Table(displayName="VoterDetailsTable", ref=f"A1:G{max_row}")
        style = TableStyleInfo(
            name="TableStyleMedium9", 
            showFirstColumn=False,
            showLastColumn=False, 
            showRowStripes=True, 
            showColumnStripes=False
        )
        tab.tableStyleInfo = style
        ws.add_table(tab)

def freeze_header(ws):
    """
    Helper: Freezes the top header row of the worksheet.
    """
    ws.freeze_panes = 'A2'

def save_workbook(wb, output_path: str):
    """
    Helper: Saves the Excel workbook.
    """
    wb.save(output_path)

def export_to_excel(raw_voters: List[Dict[str, Any]], source_pdf_name: str) -> Dict[str, Any]:
    """
    Exports structured voter records to a professionally styled Excel file.
    Includes data validation, duplicate removal, styling, and statistics reporting.
    """
    logger.info("Start Excel Export")
    start_time = time.time()
    
    # 1. Setup output directory & filename incrementing logic
    output_dir = os.path.join(os.getcwd(), "output")
    os.makedirs(output_dir, exist_ok=True)
    
    # Deriving base filename from source PDF filename
    base_name = os.path.splitext(os.path.basename(source_pdf_name))[0]
    output_filename = f"{base_name}_voter_details.xlsx"
    output_path = os.path.join(output_dir, output_filename)
    
    # Avoid overwriting existing files
    if os.path.exists(output_path):
        counter = 1
        while True:
            new_filename = f"{base_name}_voter_details_{counter}.xlsx"
            new_path = os.path.join(output_dir, new_filename)
            if not os.path.exists(new_path):
                output_path = new_path
                output_filename = new_filename
                break
            counter += 1
            
    # 2. Validate data
    valid_voters = []
    seen_epics = set()
    duplicates_count = 0
    invalid_count = 0
    
    for voter in raw_voters:
        # Skip empty records
        if not voter or not any(voter.values()):
            invalid_count += 1
            continue
            
        s_no = voter.get("S.No")
        epic = voter.get("VoterID Number")
        name = voter.get("Elector's Name")
        age = voter.get("Age")
        
        if not s_no and not epic and not name:
            invalid_count += 1
            continue
            
        # Ensure S.No is numeric
        try:
            s_no_val = int(float(str(s_no).strip()))
        except (ValueError, TypeError):
            invalid_count += 1
            continue
            
        # Ensure Age is numeric
        try:
            age_val = int(float(str(age).strip()))
        except (ValueError, TypeError):
            invalid_count += 1
            continue
            
        # Skip duplicate EPIC IDs
        if epic:
            epic_clean = str(epic).strip().upper()
            if epic_clean in seen_epics:
                duplicates_count += 1
                continue
        else:
            invalid_count += 1
            continue
            
        # Cleaned valid voter object
        voter_clean = voter.copy()
        voter_clean["S.No"] = s_no_val
        voter_clean["Age"] = age_val
        voter_clean["VoterID Number"] = epic_clean
        
        seen_epics.add(epic_clean)
        valid_voters.append(voter_clean)
        
    # 3. Create Workbook and style sheet
    wb = create_workbook()
    ws = wb.active
    ws.title = "Voter Details"
    ws.views.sheetView[0].showGridLines = True
    
    columns = [
        "S.No",
        "VoterID Number",
        "Elector's Name",
        "Guardian Name",
        "House Number",
        "Sex",
        "Age"
    ]
    
    # Apply stylings using helper functions
    style_header(ws, columns)
    style_rows(ws, start_row=2, voters=valid_voters)
    
    max_row = len(valid_voters) + 1
    create_table(ws, max_row)
    
    # Auto-filter range explicitly
    ws.auto_filter.ref = f"A1:G{max_row}"
    
    freeze_header(ws)
    auto_fit_columns(ws, min_width=15)
    
    # 4. Save Excel Workbook
    save_workbook(wb, output_path)
    
    export_time = time.time() - start_time
    
    # Logging requirements
    logger.info(f"Rows Written: {len(valid_voters)}")
    logger.info(f"Duplicates Removed: {duplicates_count}")
    logger.info(f"Export Time: {export_time:.4f}s")
    logger.info(f"Output File: {output_path}")
    
    # Relative path output format as requested: "output/AP800011_voter_details.xlsx"
    relative_file_path = f"output/{output_filename}"
    
    return {
        "status": "success",
        "file": relative_file_path,
        "records": len(valid_voters),
        "duplicates_removed": duplicates_count,
        "invalid_records": invalid_count
    }
